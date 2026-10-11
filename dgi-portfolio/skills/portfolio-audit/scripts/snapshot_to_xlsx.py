#!/usr/bin/env python3
"""Turn a live-sheet snapshot (snapshot_fetch.py) into an .xlsx the audit reads.

`audit_formulas.py` opens a workbook with `openpyxl.load_workbook(data_only=False)`,
so every cell it sees is either a formula string starting with "=" or a literal.
That is exactly the pair the bridge returns, so the conversion is direct:

    cell = formula if the formula is non-empty else the raw (unformatted) value

Using *raw* rather than *display* values matters. Display values are formatted
strings -- "$147.61", "0.82%", "TRUE" -- and the audit's type checks would read
every number as text and report the whole sheet as broken.

This builds a NEW workbook with openpyxl. It never opens an existing .xlsx, so
the project's standing rule (openpyxl load-and-save destroys cached values and
internal parts) does not apply here.

The output is a faithful stand-in for the live sheet in one respect only --
formulas and literals. It deliberately carries no formatting, no charts, no
conditional rules and no cached results, because the audit inspects none of
those. Do not hand this file to anyone as a portfolio export.

USAGE
-----
    python3 snapshot_to_xlsx.py sheet-snapshot.json -o live-audit.xlsx
    python3 audit_formulas.py live-audit.xlsx --config dgi-portfolio/config/portfolio.json
"""

import argparse
import json
import sys

import openpyxl


def cell_content(formula, raw):
    """Pick what a single cell should hold in the rebuilt workbook."""
    if isinstance(formula, str) and formula.strip():
        # Google keeps trailing newlines inside formula text; Excel does not
        # care, but they make audit findings ugly to read.
        return formula.strip()
    if raw == "" or raw is None:
        return None
    return raw


def at(grid, row_index, col_index):
    """Tolerate ragged rows -- the bridge trims trailing empties on some rows."""
    if row_index >= len(grid):
        return ""
    row = grid[row_index]
    if col_index >= len(row):
        return ""
    return row[col_index]


def build(snapshot, out_path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    stats = []
    for name, sheet in snapshot["sheets"].items():
        target = workbook.create_sheet(title=name[:31])
        formulas = sheet.get("formulas") or []
        raws = sheet.get("rawValues") or []

        height = max(len(formulas), len(raws))
        width = int(sheet.get("columns") or 0)
        written = 0

        for r in range(height):
            for c in range(width):
                content = cell_content(at(formulas, r, c), at(raws, r, c))
                if content is None:
                    continue
                target.cell(row=r + 1, column=c + 1, value=content)
                written += 1

        if sheet.get("hidden"):
            target.sheet_state = "hidden"
        stats.append((name, height, width, written))

    workbook.save(out_path)
    return stats


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("snapshot", help="JSON produced by snapshot_fetch.py")
    parser.add_argument("-o", "--out", default="live-audit.xlsx")
    args = parser.parse_args()

    with open(args.snapshot, encoding="utf-8") as fh:
        snapshot = json.load(fh)

    if "sheets" not in snapshot:
        sys.exit("%s does not look like a snapshot: no 'sheets' key." % args.snapshot)

    stats = build(snapshot, args.out)

    print("built %s from a snapshot taken %s"
          % (args.out, snapshot.get("fetched_at", "at an unrecorded time")))
    for name, rows, cols, written in stats:
        print("  %-32s %4d x %-3d  %5d non-empty cells" % (name, rows, cols, written))


if __name__ == "__main__":
    main()
