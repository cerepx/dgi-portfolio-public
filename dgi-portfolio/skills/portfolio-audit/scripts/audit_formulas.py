#!/usr/bin/env python3
"""
Deterministic integrity validator for a dividend-growth portfolio spreadsheet.

No interpretation, no guessing, no LLM. Parses formulas and reports only what it
can prove. Every finding carries a severity, a cell, evidence, and a fix.

Usage:
    python3 audit_formulas.py <file.xlsx> [--config config/portfolio.json] [--json]

Exit code 0 if no CRITICAL findings, 1 otherwise.
"""

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl missing. Run: pip install openpyxl --break-system-packages")

# --------------------------------------------------------------------------- #
CRITICAL = "CRITICAL"   # a reported number is wrong right now
SERIOUS = "SERIOUS"     # will break, or actively misleads
WARNING = "WARNING"     # fragile, needs manual verification
INFO = "INFO"           # observation, no action

SEVERITY_ORDER = {CRITICAL: 0, SERIOUS: 1, WARNING: 2, INFO: 3}

GOOGLE_ONLY_FUNCTIONS = [
    "GOOGLEFINANCE", "SPARKLINE", "IMPORTHTML", "IMPORTXML",
    "IMPORTDATA", "IMPORTRANGE", "IMAGE", "QUERY", "FLATTEN",
]

RE_APPS_SCRIPT = re.compile(r"=\s*(get[A-Z]\w*)\s*\(")
RE_DUMMY = re.compile(r"__xludf\.DUMMYFUNCTION")
RE_SUM_RANGE = re.compile(r"SUM\(\s*([A-Z]{1,3})(\d+)\s*:\s*([A-Z]{1,3})(\d+)\s*\)", re.I)
RE_CELL_ERROR = re.compile(r"#(REF!|DIV/0!|VALUE!|NAME\?|NUM!|N/A|NULL!|SPILL!|CALC!)")
RE_FREQUENCY = re.compile(r"\*(\d+)\*")
RE_LEADING_FACTOR = re.compile(r"^=\s*(\d*\.?\d+)\s*\*\s*\(")


class Report:
    def __init__(self):
        self.findings = []

    def add(self, severity, sheet, cell, title, evidence, fix):
        self.findings.append({
            "severity": severity,
            "sheet": sheet,
            "cell": cell or "",
            "title": title,
            "evidence": evidence,
            "fix": fix,
        })

    def sorted_findings(self):
        return sorted(self.findings,
                      key=lambda f: (SEVERITY_ORDER[f["severity"]], f["sheet"], f["cell"]))

    def count(self, severity):
        return sum(1 for f in self.findings if f["severity"] == severity)


def cell_value(cell):
    return cell.value if cell is not None else None


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def flatten(formula):
    """Strip the newlines Google Sheets injects inside formulas."""
    return re.sub(r"\s+", "", formula) if isinstance(formula, str) else formula


def block_rows(spec, name):
    block = spec.get("blocks", {}).get(name, {})
    first, last = block.get("first_row"), block.get("last_row")
    if first is None or last is None:
        return []
    return list(range(first, last + 1))


def column_by_meaning(columns, meaning, default=None):
    return next((k for k, v in columns.items() if v.get("meaning") == meaning), default)


def tickers_in_rows(sheet, ticker_col, rows):
    out = {}
    for row in rows:
        value = cell_value(sheet[f"{ticker_col}{row}"])
        if value not in (None, ""):
            out[row] = str(value).strip()
    return out


# --------------------------------------------------------------------------- #
# Per-payment dividend sheet. Several checks change behaviour entirely
# depending on whether one exists, because it turns assumptions the older
# checks were written to flag into facts they can verify.


def payment_sheet(workbook, config):
    """Return (name, spec) of the configured per-payment sheet, or (None, None)."""
    for name, spec in config.get("sheets", {}).items():
        if spec.get("role") == "dividend_payments" and name in workbook.sheetnames:
            return name, spec
    return None, None


def read_payments(sheet, spec):
    """Every payment row as a dict. Rows without a ticker are padding."""
    columns = spec.get("columns", {})
    col_year = column_by_meaning(columns, "payment_year", "A")
    col_month = column_by_meaning(columns, "payment_month", "B")
    col_ticker = column_by_meaning(columns, "ticker", "C")
    col_type = column_by_meaning(columns, "payment_type", "I")

    payments = []
    for row in range(spec.get("first_data_row", 2), sheet.max_row + 1):
        ticker = cell_value(sheet[f"{col_ticker}{row}"])
        if ticker in (None, ""):
            continue

        def number(col):
            value = cell_value(sheet[f"{col}{row}"])
            try:
                return int(float(value))
            except (TypeError, ValueError):
                return None

        kind = cell_value(sheet[f"{col_type}{row}"])
        payments.append({
            "row": row,
            "ticker": str(ticker).strip(),
            "year": number(col_year),
            "month": number(col_month),
            "type": str(kind).strip().lower() if kind not in (None, "") else "",
        })
    return payments


def trailing_year_counts(payments, kind="regular"):
    """Regular payments per ticker over EACH TICKER'S OWN last 12 months.

    **The window is per ticker, not per sheet.** BACKLOG item 21.

    It used to be anchored on the newest payment anywhere in the sheet, and the
    reason was sound: anchoring on the newest row rather than on today stops the
    count shrinking merely because the sheet has not been updated this month.
    The mirror failure was not anticipated - **one ticker paying early in the
    month advances the window past another ticker's payment that is not due
    yet.**

    Measured 2026-09-05. Adding two real September-1 payments (V, ZTS) moved the
    window to 2025-10..2026-09, which dropped the September 2025 payment of
    QCOM, O, JNJ, NEE and ADM while their September 2026 payment had not
    happened. Five correct positions came out as SERIOUS, which this audit
    defines as "annual income is wrong by the ratio of the two numbers". It was
    not wrong at all.

    Returns `(counts, first_seen, windows)`, where `windows` maps each ticker to
    its own `(oldest, newest)` month stamp.
    """
    newest_by_ticker, first_seen = {}, {}
    for payment in payments:
        if not payment["year"] or not payment["month"]:
            continue
        ticker = payment["ticker"]
        stamp = payment["year"] * 12 + payment["month"]
        newest_by_ticker[ticker] = max(newest_by_ticker.get(ticker, stamp), stamp)
        first_seen[ticker] = min(first_seen.get(ticker, stamp), stamp)

    windows = {t: (newest - 11, newest) for t, newest in newest_by_ticker.items()}

    counts = defaultdict(int)
    for payment in payments:
        if not payment["year"] or not payment["month"]:
            continue
        if payment["type"] != kind:
            continue
        window = windows.get(payment["ticker"])
        stamp = payment["year"] * 12 + payment["month"]
        if window and window[0] <= stamp <= window[1]:
            counts[payment["ticker"]] += 1
    return counts, first_seen, windows


# --------------------------------------------------------------------------- #
# CHECK 1 - dead Google formulas
def check_dead_google_formulas(workbook, config, report):
    dummy_per_sheet = defaultdict(int)
    apps_script = defaultdict(list)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell_value(cell)
                if not isinstance(value, str):
                    continue
                if RE_DUMMY.search(value):
                    dummy_per_sheet[sheet.title] += 1
                match = RE_APPS_SCRIPT.search(value)
                if match:
                    apps_script[match.group(1)].append(f"{sheet.title}!{cell.coordinate}")

    total = sum(dummy_per_sheet.values())
    if total:
        breakdown = ", ".join(f"{k}: {v}" for k, v in
                              sorted(dummy_per_sheet.items(), key=lambda x: -x[1]))
        report.add(
            CRITICAL, "(whole file)", "",
            f"{total} dead Google formulas (__xludf.DUMMYFUNCTION)",
            f"Breakdown by sheet - {breakdown}. Affected functions: "
            f"{', '.join(GOOGLE_ONLY_FUNCTIONS)}. Every one of those cells displays a "
            f"value CACHED at export time, not live data.",
            "This file is a dead export. Do not treat it as the source of truth and do "
            "not write to it - re-uploading to Google Sheets permanently destroys the "
            "original formulas. Work directly in Google Sheets via the Apps Script web app.",
        )

    for function_name, cells in apps_script.items():
        shown = ", ".join(cells[:5]) + ("..." if len(cells) > 5 else "")
        report.add(
            SERIOUS, "(whole file)", shown,
            f"Apps Script function '{function_name}()' in {len(cells)} cell(s)",
            f"'{function_name}' is defined in the Google Sheets script editor. "
            f"In Excel or LibreOffice it returns #NAME?.",
            f"Keep the source of {function_name}() in a versioned file. If the site it "
            f"scrapes changes structure, the function silently returns the last good "
            f"value with no visible error and the price freezes. Add a last-updated "
            f"timestamp next to it so staleness is detectable.",
        )


# --------------------------------------------------------------------------- #
# CHECK 2 - total row range consistency
def check_total_ranges(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        total_rows = spec.get("total_rows") or []
        if not total_rows:
            continue

        sheet = workbook[sheet_name]
        columns = spec.get("columns", {})
        ticker_col = column_by_meaning(columns, "ticker", "B")
        dgi = block_rows(spec, "dgi")
        non_dgi = block_rows(spec, "non_dgi")
        if not dgi:
            continue

        dgi_tickers = tickers_in_rows(sheet, ticker_col, dgi)
        last_dgi_row = max(dgi_tickers) if dgi_tickers else spec["blocks"]["dgi"]["first_row"]
        non_dgi_tickers = tickers_in_rows(sheet, ticker_col, non_dgi)

        for total_row in total_rows:
            sums = {}
            for cell in sheet[total_row]:
                value = cell_value(cell)
                if not is_formula(value):
                    continue
                match = RE_SUM_RANGE.search(flatten(value))
                if match:
                    sums[cell.coordinate] = (int(match.group(2)), int(match.group(4)))
            if not sums:
                continue

            # 2a - different end rows across aggregated columns
            by_end = defaultdict(list)
            for coord, (_, end) in sums.items():
                by_end[end].append(coord)
            if len(by_end) > 1:
                desc = "; ".join(f"ends at row {end}: {', '.join(sorted(c))}"
                                 for end, c in sorted(by_end.items()))
                report.add(
                    CRITICAL, sheet_name, f"row {total_row}",
                    "Total row sums ranges of different lengths",
                    f"{desc}. Columns describing the same positions cover different "
                    f"rows, so the totals are not comparable with each other.",
                    f"Align every SUM to the same range.",
                )

            # 2b - DGI positions left out of the aggregate
            for coord, (start, end) in sorted(sums.items()):
                if end < last_dgi_row:
                    missing = [f"{t} (row {r})" for r, t in sorted(dgi_tickers.items())
                               if r > end]
                    if missing:
                        report.add(
                            CRITICAL, sheet_name, coord,
                            "DGI positions excluded from the total",
                            f"{coord} = SUM(...{end}) but DGI data runs to row "
                            f"{last_dgi_row}. Excluded: {', '.join(missing)}.",
                            f"Extend the range to at least row {last_dgi_row}.",
                        )

                # 2c - non-DGI rows swallowed by an aggregate that must exclude them
                swallowed = [f"{t} (row {r})" for r, t in sorted(non_dgi_tickers.items())
                             if start <= r <= end]
                col_letter = re.match(r"([A-Z]+)", coord).group(1)
                meta = columns.get(col_letter, {})
                if swallowed and meta.get("dgi_only_aggregate"):
                    report.add(
                        CRITICAL, sheet_name, coord,
                        f"Non-DGI positions included in the DGI aggregate "
                        f"'{meta.get('name', col_letter)}'",
                        f"{coord} = SUM({col_letter}{start}:{col_letter}{end}) covers "
                        f"non-DGI rows: {', '.join(swallowed)}. This column feeds DGI "
                        f"reporting and must exclude tracking-only positions.",
                        f"Restrict the range to the DGI block "
                        f"({col_letter}{spec['blocks']['dgi']['first_row']}:"
                        f"{col_letter}{spec['blocks']['dgi']['last_row'] + 1}), or use "
                        f"SUMIFS with an explicit 'DGI' flag column so the boundary is "
                        f"data, not a hardcoded row number.",
                    )

            # 2d - stray values in the total row
            for cell in sheet[total_row]:
                value = cell_value(cell)
                if value in (None, "") or is_formula(value):
                    continue
                if isinstance(value, str) and len(value) > 3 and " " in value:
                    continue  # legitimate label
                report.add(
                    SERIOUS, sheet_name, cell.coordinate,
                    "Stray value in the totals row",
                    f"{cell.coordinate} = {value!r} - neither a formula nor a label. "
                    f"Looks like a forgotten note or an accidental overwrite.",
                    "Delete it, or move it to the Watchlist or a dedicated notes column.",
                )


# --------------------------------------------------------------------------- #
# CHECK 3 - the separator label that documents the non-DGI block
def check_non_dgi_separator(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        separator = spec.get("blocks", {}).get("separator")
        if not separator:
            continue
        sheet = workbook[sheet_name]
        cell = separator.get("cell")
        expected = separator.get("expected_label", "")
        actual = cell_value(sheet[cell])
        if actual is None or expected.lower() not in str(actual).lower():
            report.add(
                SERIOUS, sheet_name, cell,
                "Non-DGI block separator label missing or changed",
                f"{cell} = {actual!r}, expected something containing {expected!r}. "
                f"The label is the only in-sheet documentation that rows "
                f"{spec['blocks']['non_dgi']['first_row']}-"
                f"{spec['blocks']['non_dgi']['last_row']} are tracking-only.",
                "Restore the label. Better still, add an explicit 'DGI' TRUE/FALSE "
                "column so the distinction is data rather than a text note plus "
                "hardcoded row numbers.",
            )


# --------------------------------------------------------------------------- #
# CHECK 4 - non-DGI rows must not carry DGI metrics
def check_non_dgi_isolation(workbook, config, report):
    marker = None
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        non_dgi = block_rows(spec, "non_dgi")
        if not non_dgi:
            continue
        sheet = workbook[sheet_name]
        columns = spec.get("columns", {})
        marker = spec.get("unavailable_marker", "N/A")
        ticker_col = column_by_meaning(columns, "ticker", "B")
        dgi_only_cols = [k for k, v in columns.items() if v.get("dgi_only")]

        populated = []
        for row in non_dgi:
            ticker = cell_value(sheet[f"{ticker_col}{row}"])
            if not ticker:
                continue
            for col in dgi_only_cols:
                value = cell_value(sheet[f"{col}{row}"])
                if value in (None, "", marker):
                    continue
                if isinstance(value, (int, float)) and value == 0:
                    continue
                populated.append(f"{ticker} {col}{row} = {str(value)[:40]}")

        if populated:
            report.add(
                SERIOUS, sheet_name, ", ".join(p.split()[1] for p in populated[:6]),
                f"DGI metrics populated on {len(populated)} non-DGI cell(s)",
                f"Tracking-only positions carry values in DGI-only columns: "
                f"{'; '.join(populated[:6])}"
                f"{f' and {len(populated) - 6} more' if len(populated) > 6 else ''}. "
                f"These feed dividend yield, income weight and growth reporting.",
                f"Set them to '{marker}' so they cannot leak into DGI aggregates.",
            )
        else:
            report.add(
                INFO, sheet_name, "",
                "Non-DGI block correctly isolated",
                f"Rows {min(non_dgi)}-{max(non_dgi)} carry '{marker}' in every DGI-only "
                f"column. Empty dividend cells there are intentional, not a data gap.",
                "No action needed.",
            )


# --------------------------------------------------------------------------- #
# CHECK 5 - weight denominators
def check_weight_denominators(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        columns = spec.get("columns", {})
        weight_cols = [k for k, v in columns.items() if v.get("expected_sum") == 1.0]
        if not weight_cols:
            continue
        sheet = workbook[sheet_name]
        dgi = block_rows(spec, "dgi")
        non_dgi = block_rows(spec, "non_dgi")
        all_rows = dgi + non_dgi
        if not all_rows:
            continue

        for col in weight_cols:
            shapes = defaultdict(list)
            for row in all_rows:
                value = cell_value(sheet[f"{col}{row}"])
                if not is_formula(value):
                    continue
                flat = flatten(value)
                if "/" not in flat:
                    continue
                denominator = flat.split("/", 1)[1]
                shape = re.sub(r"[A-Z]{1,3}(?=\d)", "@", denominator)
                shapes[shape].append(f"{col}{row}")

            if len(shapes) > 1:
                desc = "; ".join(
                    f"'{shape}' in {', '.join(cells[:4])}"
                    f"{'...' if len(cells) > 4 else ''}"
                    for shape, cells in shapes.items())
                report.add(
                    CRITICAL, sheet_name, f"column {col}",
                    f"Column '{columns[col]['name']}' mixes {len(shapes)} denominator shapes",
                    f"{desc}. Weights computed on different bases do not sum to 100% "
                    f"and cannot be compared with each other.",
                    "Pick one denominator for the whole column. If DGI and non-DGI need "
                    "different bases, use two columns - 'Weight (DGI)' and 'Weight "
                    "(total)' - not two formula shapes in one column.",
                )

            # non-DGI weights measured against a DGI-only denominator
            total_rows = spec.get("total_rows") or []
            if not total_rows or not non_dgi:
                continue
            value_col = column_by_meaning(columns, "current_value", "F")
            total_formula = cell_value(sheet[f"{value_col}{total_rows[0]}"])
            if not is_formula(total_formula):
                continue
            match = RE_SUM_RANGE.search(flatten(total_formula))
            if not match:
                continue
            total_end = int(match.group(4))
            offenders = []
            for row in non_dgi:
                value = cell_value(sheet[f"{col}{row}"])
                if not is_formula(value):
                    continue
                flat = flatten(value)
                if f"${value_col}${total_rows[0]}" in flat and row > total_end \
                        and f"+{value_col}{row}" not in flat and f"+{value_col}{row}/" not in flat:
                    offenders.append(f"{col}{row}")
            if offenders:
                report.add(
                    CRITICAL, sheet_name, ", ".join(offenders),
                    "Non-DGI weights divided by a total that excludes them",
                    f"{', '.join(offenders)} divide by "
                    f"${value_col}${total_rows[0]} = SUM(...{total_end}), which stops "
                    f"before the non-DGI block. The position is not in its own "
                    f"denominator, so its weight is overstated and the column cannot "
                    f"sum to 100%.",
                    f"Add a real grand total that covers every position "
                    f"(DGI plus non-DGI) and divide by that. Keep the DGI-only total as "
                    f"a separate cell for the dividend reporting.",
                )


# --------------------------------------------------------------------------- #
# CHECK 6 - dividend tax factor vs the year the payment was received
def check_tax_factors(workbook, config, report):
    tax = config.get("dividend_tax", {})
    legacy = tax.get("legacy_net_factor")
    current = tax.get("current_net_factor")
    current_year = tax.get("current_year")
    if legacy is None or current is None:
        return

    known_factors = {b["net_factor"]: b for b in tax.get("brackets", [])}

    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        columns = spec.get("columns", {})
        col = column_by_meaning(columns, "dividends_received_net")
        if not col:
            continue
        sheet = workbook[sheet_name]
        ticker_col = column_by_meaning(columns, "ticker", "B")
        rows = block_rows(spec, "dgi")

        legacy_cells, unknown_cells = [], []
        for row in rows:
            value = cell_value(sheet[f"{col}{row}"])
            if not is_formula(value):
                continue
            flat = flatten(value)
            match = RE_LEADING_FACTOR.match(flat)
            if not match:
                continue
            factor = float(match.group(1))
            ticker = cell_value(sheet[f"{ticker_col}{row}"]) or "?"
            if abs(factor - legacy) < 1e-9:
                legacy_cells.append(f"{ticker} ({col}{row})")
            elif factor not in known_factors:
                unknown_cells.append(f"{ticker} ({col}{row}) factor {factor}")

        if legacy_cells:
            report.add(
                CRITICAL, sheet_name, f"column {col}",
                f"Dividend tax factor {legacy} applied to all payments, including "
                f"{current_year} income taxed at {round((1 - current) * 100)}%",
                f"{len(legacy_cells)} cells still multiply the entire cumulative total "
                f"by {legacy}: {', '.join(legacy_cells[:6])}"
                f"{f' and {len(legacy_cells) - 6} more' if len(legacy_cells) > 6 else ''}. "
                f"The Romanian dividend tax rose from "
                f"{round((1 - legacy) * 100)}% to {round((1 - current) * 100)}% on "
                f"1 January {current_year}, so every payment received in "
                f"{current_year} or later is overstated by roughly "
                f"{(current / legacy - 1) * -100:.1f}%.",
                f"Do NOT simply replace {legacy} with {current} - that would "
                f"retroactively re-tax pre-{current_year} income and understate what you "
                f"actually received. The factor must be applied per payment by year of "
                f"receipt: {legacy}*(pre-{current_year} payments) + "
                f"{current}*({current_year}+ payments). A flat string of summed numbers "
                f"carries no dates, so it cannot be split correctly - split column {col} "
                f"into a per-payment sheet first.",
            )

        for entry in unknown_cells:
            report.add(
                WARNING, sheet_name, col,
                "Unrecognised tax factor",
                f"{entry} does not match any configured bracket "
                f"({', '.join(str(k) for k in known_factors)}).",
                "Check it against the rate in force for the year of receipt.",
            )


# --------------------------------------------------------------------------- #
# CHECK 7 - non-auditable dividend history
def check_dividend_history_shape(workbook, config, report):
    # If a per-payment sheet exists, this whole check is obsolete: the column
    # it inspects is a SUMIFS into that sheet, not a flat string of summed
    # amounts. Parsing a SUMIFS for "payment magnitudes" finds the digits of
    # cell references and reports them as payments -- which is exactly what it
    # did on 2026-08-01, claiming "17 payments recorded as inline sums" for 17
    # rows that each contained one SUMIFS.
    payments_sheet, _ = payment_sheet(workbook, config)

    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        columns = spec.get("columns", {})
        col = column_by_meaning(columns, "dividends_received_net")
        if not col:
            continue
        sheet = workbook[sheet_name]

        if payments_sheet:
            report.add(
                INFO, sheet_name, f"column {col}",
                "Dividend history is per-payment, not inline sums",
                f"Column {col} aggregates '{payments_sheet}', which stores one row "
                f"per payment with its own date, share count and tax factor. The "
                f"inline-sum shape this check was written for no longer exists.",
                "No action needed. The integrity of the payment rows themselves is "
                "covered by the per-payment sheet check.",
            )
            continue
        ticker_col = column_by_meaning(columns, "ticker", "B")
        inconsistent, total_payments = [], 0

        for row in block_rows(spec, "dgi"):
            value = cell_value(sheet[f"{col}{row}"])
            if not is_formula(value) or "(" not in str(value):
                continue
            flat = flatten(value)
            numbers = [float(x) for x in
                       re.findall(r"\d+\.?\d*", flat.split("(", 1)[1])]
            numbers = [n for n in numbers if n > 0]
            total_payments += len(numbers)
            if len(numbers) < 4:
                continue
            median = sorted(numbers)[len(numbers) // 2]
            if median <= 0:
                continue
            if any(n > 2.5 * median or n < 0.4 * median for n in numbers):
                ticker = cell_value(sheet[f"{ticker_col}{row}"]) or "?"
                inconsistent.append(
                    f"{ticker} ({col}{row}): payments from {min(numbers)} to "
                    f"{max(numbers)}, median {median}")

        if inconsistent:
            report.add(
                SERIOUS, sheet_name, f"column {col}",
                f"Dividend history not interpretable for {len(inconsistent)} position(s)",
                f"Payments are stored as a flat string of summed amounts with no date and "
                f"no share count at the time of payment. Widely different magnitudes in "
                f"the same string: {'; '.join(inconsistent[:5])}"
                f"{f' and {len(inconsistent) - 5} more' if len(inconsistent) > 5 else ''}. "
                f"It is impossible to tell whether the variation comes from a special "
                f"dividend, a raise, or buying more shares between payments. That is also "
                f"why DGR cannot be computed from your own data, and why the tax factor "
                f"cannot be split by year.",
                f"Split into a 'Dividends received' sheet, one row per payment: "
                f"ex-date | pay date | ticker | gross per share | shares | gross | "
                f"tax rate | net. Column {col} then becomes a SUMIFS on ticker. This one "
                f"change unlocks real DGR, correct per-year taxation, and Declaratia "
                f"Unica figures straight from the sheet.",
            )
        elif total_payments:
            report.add(
                INFO, sheet_name, f"column {col}",
                f"{total_payments} payments recorded as inline sums",
                "No magnitude outliers detected, but the format still stores no dates.",
                "Per-payment rows remain the prerequisite for DGR and per-year tax.",
            )


# --------------------------------------------------------------------------- #
# CHECK 8 - hardcoded payment frequency
def check_hardcoded_frequency(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames or spec.get("role") != "open_positions":
            continue
        columns = spec.get("columns", {})
        income_col = column_by_meaning(columns, "annual_income")
        if not income_col:
            continue
        sheet = workbook[sheet_name]
        ticker_col = column_by_meaning(columns, "ticker", "B")
        hardcoded = []
        for row in block_rows(spec, "dgi"):
            value = cell_value(sheet[f"{income_col}{row}"])
            if not is_formula(value):
                continue
            match = RE_FREQUENCY.search(flatten(value))
            if not match:
                continue
            ticker = cell_value(sheet[f"{ticker_col}{row}"]) or "?"
            hardcoded.append((str(ticker).strip(), int(match.group(1))))

        if not hardcoded:
            continue

        # With a per-payment sheet the multiplier is no longer an assumption to
        # flag -- it is a claim to test. Only disagreement with the real
        # payment history is worth a warning; blanket-warning on all of them
        # buries a genuine mismatch in noise, which is what this check did
        # until 2026-08-01.
        payments_name, payments_spec = payment_sheet(workbook, config)
        if payments_name:
            payments = read_payments(workbook[payments_name], payments_spec)
            counts, first_seen, windows = trailing_year_counts(payments)

            mismatched, unverifiable = [], []
            for ticker, frequency in hardcoded:
                observed = counts.get(ticker, 0)
                started = first_seen.get(ticker)
                window = windows.get(ticker)
                full_year = started is not None and window and \
                    started <= window[0]
                if not full_year:
                    unverifiable.append(f"{ticker} ({observed} payment(s) on record)")
                elif observed != frequency:
                    mismatched.append(
                        f"{ticker}: formula says {frequency}x, history shows "
                        f"{observed} regular payment(s)")

            if mismatched:
                report.add(
                    SERIOUS, sheet_name, f"column {income_col}",
                    f"Payment frequency disagrees with payment history for "
                    f"{len(mismatched)} position(s)",
                    f"Counted against '{payments_name}' over each ticker's own "
                    f"trailing twelve months: {'; '.join(mismatched)}. Annual "
                    f"income is wrong by the "
                    f"ratio of the two numbers.",
                    f"Correct the multiplier in {income_col}, or confirm the payment "
                    f"history is complete for that ticker. Do not assume the formula "
                    f"is right because it always was.",
                )
            else:
                verified = len(hardcoded) - len(unverifiable)
                report.add(
                    INFO, sheet_name, f"column {income_col}",
                    f"Payment frequency verified for {verified} of "
                    f"{len(hardcoded)} position(s)",
                    f"Every multiplier with a full year of history in "
                    f"'{payments_name}' matches the real count of regular payments."
                    + (f" Not yet verifiable, too little history: "
                       f"{', '.join(unverifiable)}." if unverifiable else ""),
                    "No action needed. The remainder resolve themselves as payments "
                    "accumulate; a counted frequency would be worse, since one "
                    "payment on record would imply an annual frequency of 1.",
                )
            continue

        report.add(
            WARNING, sheet_name, f"column {income_col}",
            f"Payment frequency hardcoded for {len(hardcoded)} position(s)",
            f"Each formula is dividend x frequency x shares with the frequency "
            f"written into the formula: "
            f"{', '.join(f'{t} {f}x' for t, f in hardcoded[:8])}"
            f"{f' and {len(hardcoded) - 8} more' if len(hardcoded) > 8 else ''}. "
            f"There is no per-payment sheet to check it against, so none of it can "
            f"be verified. A missed raise understates annual income; a missed cut "
            f"overstates it - and a cut is the only sell signal in the DGI strategy.",
            "Build a per-payment 'Dividends received' sheet and declare it in the "
            "config with role 'dividend_payments'. This check then verifies every "
            "multiplier instead of merely suspecting it.",
        )


# --------------------------------------------------------------------------- #
# CHECK 9 - mislabeled columns
def check_labels(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        header_row = spec.get("header_row", 1)
        for col, meta in spec.get("columns", {}).items():
            if not meta.get("mislabeled"):
                continue
            report.add(
                SERIOUS, sheet_name, f"{col}{header_row}",
                f"Label '{meta['name']}' contradicts the formula",
                meta.get("note", ""),
                "Rename the column to what it actually computes (Yield on Cost) and add "
                "a separate column for true Dividend Yield = annual income / current "
                "value. YoC measures how your invested money is performing; current "
                "yield tells you whether the stock is cheap today. A 6% YoC on a stock "
                "yielding 2.5% is not a buying opportunity, but the wrong label suggests "
                "it is.",
            )


# --------------------------------------------------------------------------- #
# CHECK 10 - duplicate and malformed tickers
def check_tickers(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        columns = spec.get("columns", {})
        col = column_by_meaning(columns, "ticker")
        if not col:
            continue
        sheet = workbook[sheet_name]
        rows = block_rows(spec, "dgi") + block_rows(spec, "non_dgi")
        if not rows:
            rows = range(spec.get("first_data_row", 2), sheet.max_row + 1)

        seen = defaultdict(list)
        for row in rows:
            value = cell_value(sheet[f"{col}{row}"])
            if value in (None, ""):
                continue
            seen[str(value).strip().upper()].append(row)
            if isinstance(value, str) and value != value.strip():
                report.add(
                    WARNING, sheet_name, f"{col}{row}",
                    "Ticker has surrounding whitespace",
                    f"{col}{row} = {value!r}. VLOOKUP and API calls fail on this.",
                    "Apply =UPPER(TRIM(...)) or clean it manually.",
                )

        # A repeated ticker is a duplicated position only where each row IS a
        # position. On a per-payment sheet repetition is the entire design -
        # nine KO rows are nine dividends, not nine holdings. Reporting those
        # as CRITICAL produced 12 false alarms the moment the payment sheet was
        # added to the config on 2026-08-01.
        if spec.get("role") != "open_positions":
            continue

        for ticker, rows_found in seen.items():
            if len(rows_found) > 1:
                report.add(
                    CRITICAL, sheet_name,
                    ", ".join(f"{col}{r}" for r in rows_found),
                    f"Duplicate ticker: {ticker}",
                    f"{ticker} appears on rows {', '.join(map(str, rows_found))}. "
                    f"The position is counted more than once in totals and weights.",
                    "Consolidate into a single row with a weighted average cost.",
                )


# --------------------------------------------------------------------------- #
# CHECK 11 - sector misclassification
def check_sectors(workbook, config, report):
    expected = {k: v for k, v in config.get("expected_sectors", {}).items()
                if not k.startswith("_")}
    if not expected:
        return
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        columns = spec.get("columns", {})
        ticker_col = column_by_meaning(columns, "ticker")
        sector_col = column_by_meaning(columns, "sector")
        if not (ticker_col and sector_col):
            continue
        sheet = workbook[sheet_name]
        for row in block_rows(spec, "dgi") + block_rows(spec, "non_dgi"):
            ticker = cell_value(sheet[f"{ticker_col}{row}"])
            if not ticker:
                continue
            ticker = str(ticker).strip().upper()
            actual = cell_value(sheet[f"{sector_col}{row}"])
            want = expected.get(ticker)
            if want and actual and str(actual).strip() != want:
                report.add(
                    CRITICAL, sheet_name, f"{sector_col}{row}",
                    f"{ticker} classified in the wrong sector",
                    f"{sector_col}{row} = {str(actual)!r}, expected {want!r}. "
                    f"Sector concentration checks and any sector pie chart are wrong "
                    f"by this position's weight.",
                    f"Set it to {want!r}. Verify against the GICS sector from FMP rather "
                    f"than entering it by hand.",
                )


# --------------------------------------------------------------------------- #
# CHECK 12 - formula errors
def check_formula_errors(workbook, config, report):
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell_value(cell)
                if isinstance(value, str) and RE_CELL_ERROR.search(value) \
                        and not RE_DUMMY.search(value) and value.startswith("="):
                    report.add(
                        CRITICAL, sheet.title, cell.coordinate,
                        "Formula error",
                        f"{cell.coordinate} = {value[:120]}",
                        "Fix the reference, or wrap in IFERROR with an explicit default.",
                    )


# --------------------------------------------------------------------------- #
# CHECK 13 - fragile fixed-row references
def check_fragile_references(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames or "note" not in spec:
            continue
        report.add(
            WARNING, sheet_name, "",
            "Fragile fixed-row references",
            spec["note"],
            "Replace absolute row references with SUMIFS on a year column, or use named "
            "ranges, so inserting a row cannot silently shift the result.",
        )


# --------------------------------------------------------------------------- #
# CHECK 14 - type consistency within a column
def check_type_consistency(workbook, config, report):
    for sheet_name, spec in config.get("sheets", {}).items():
        if sheet_name not in workbook.sheetnames:
            continue
        columns = spec.get("columns", {})
        sheet = workbook[sheet_name]
        rows = block_rows(spec, "dgi")
        if not rows:
            continue
        marker = spec.get("unavailable_marker", "N/A")

        for col, meta in columns.items():
            if meta.get("type") not in ("formula", "manual_formula"):
                continue
            formulas, literals = [], []
            for row in rows:
                value = cell_value(sheet[f"{col}{row}"])
                if value in (None, "", marker):
                    continue
                (formulas if is_formula(value) else literals).append(f"{col}{row}")
            if formulas and literals:
                report.add(
                    WARNING, sheet_name, ", ".join(literals[:6]),
                    f"Column '{meta['name']}' mixes formulas and hardcoded values",
                    f"{len(formulas)} formula cells and {len(literals)} literal cells: "
                    f"{', '.join(literals[:6])}"
                    f"{f' and {len(literals) - 6} more' if len(literals) > 6 else ''}. "
                    f"Literals do not update when their inputs change.",
                    "Make the whole column consistent so a stale value is impossible.",
                )


# --------------------------------------------------------------------------- #
def format_text(report, path):
    lines = ["=" * 78,
             f"PORTFOLIO INTEGRITY AUDIT - {Path(path).name}",
             "=" * 78]
    counts = {s: report.count(s) for s in (CRITICAL, SERIOUS, WARNING, INFO)}
    lines.append(f"CRITICAL: {counts[CRITICAL]}   SERIOUS: {counts[SERIOUS]}   "
                 f"WARNING: {counts[WARNING]}   INFO: {counts[INFO]}")
    lines.append("")

    if not report.findings:
        lines.append("No issues detected.")
        return "\n".join(lines)

    current = None
    for index, finding in enumerate(report.sorted_findings(), 1):
        if finding["severity"] != current:
            current = finding["severity"]
            lines += ["", f"--- {current} " + "-" * max(0, 72 - len(current))]
        where = f"{finding['sheet']}!{finding['cell']}" if finding["cell"] else finding["sheet"]
        lines += ["",
                  f"[{index}] {finding['title']}",
                  f"     Where:    {where}",
                  f"     Evidence: {finding['evidence']}",
                  f"     Fix:      {finding['fix']}"]

    lines += ["", "=" * 78,
              "NOTHING was modified. This report is strictly read-only.",
              "=" * 78]
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# CHECK 15 - integrity of the per-payment dividend sheet
def check_payment_sheet(workbook, config, report):
    """Guard the sheet that now feeds every dividend aggregate.

    Written 2026-08-01 after the sheet spent a week with no coverage at all,
    during which an ARRAYFORMULA spill was broken by a batch of literal writes
    and only caught because a human re-read the aggregates.
    """
    sheet_name, spec = payment_sheet(workbook, config)
    if not sheet_name:
        return

    sheet = workbook[sheet_name]
    columns = spec.get("columns", {})
    first_row = spec.get("first_data_row", 2)

    # --- the spill anchors. This is the expensive failure mode. ------------- #
    # A spilled cell holds a value and no formula, which is indistinguishable
    # from a literal. The anchor is the only cell that can be judged: if it
    # still holds an ARRAYFORMULA the column below it is a spill; if it holds a
    # literal, or nothing while cells below have values, the spill is broken
    # and every row above the breakage has silently lost its value.
    for col, anchor in (spec.get("spill_anchors") or {}).items():
        anchor_value = cell_value(sheet[anchor])
        below = [row for row in range(first_row + 1, sheet.max_row + 1)
                 if cell_value(sheet[f"{col}{row}"]) not in (None, "")]

        if is_formula(anchor_value):
            continue
        if not below:
            continue

        # Two states reach this point and they need different words. BACKLOG
        # item 18: this check treats the config as ground truth and the sheet as
        # the thing under test, so the only sentence it could form was "the
        # spill is broken" - even when the truth was that the config had gone
        # stale (item 10, a false CRITICAL for two weeks).
        #
        # An EMPTY anchor with values below really is a break: a column that was
        # deliberately made manual carries a value in its first data row.
        # A LITERAL anchor is genuinely ambiguous from the sheet alone.
        #
        # Severity stays CRITICAL in both, because the cost if it is real is
        # the whole column. What changes is the remediation - and that is the
        # part that mattered. "Clear every literal below" destroys every row of
        # a deliberately manual column.
        if anchor_value in (None, ""):
            report.add(
                CRITICAL, sheet_name, anchor,
                f"ARRAYFORMULA spill broken in column {col}",
                f"{anchor} should hold the ARRAYFORMULA that fills {col}"
                f"{first_row} downwards, but is empty, while {len(below)} "
                f"cell(s) below it still carry values. A column that was made "
                f"manual on purpose would carry a value here, so this is a "
                f"real break: the spill collapsed and every row that depended "
                f"on it emptied silently.",
                f"Restore the ARRAYFORMULA in {anchor} and clear every literal "
                f"in {col}{first_row + 1} downwards. When appending payments, "
                f"write only {', '.join(spec.get('manual_columns') or [])} - "
                f"never the spilled columns.",
            )
        else:
            report.add(
                CRITICAL, sheet_name, anchor,
                f"Column {col} does not match the config - one of the two is stale",
                f"The config declares {anchor} an ARRAYFORMULA anchor, but it "
                f"holds a literal value and {len(below)} cell(s) below it carry "
                f"values. From the sheet alone these two are indistinguishable: "
                f"either the spill was overwritten and the column has silently "
                f"emptied, or the column was deliberately made manual and the "
                f"config was never told. **This check cannot tell you which.**",
                f"SETTLE IT BEFORE TOUCHING ANYTHING. Ask whether {col} was "
                f"made manual on purpose, or compare a few values against what "
                f"the ARRAYFORMULA would produce. If it was made manual, the "
                f"fix is in config/portfolio.json - move {col} out of "
                f"spill_anchors into manual_columns and update _append_rule. "
                f"DO NOT clear the column: if it is manual, every one of those "
                f"{len(below)} rows is real data and nothing will refill them.",
            )

    # --- literals written into a spilled column ---------------------------- #
    # Distinct from the above: the anchor survives, but someone typed into a
    # cell below it. Google keeps the literal and the spill stops there.
    for col in (spec.get("spill_anchors") or {}):
        anchor_value = cell_value(sheet[spec["spill_anchors"][col]])
        if not is_formula(anchor_value):
            continue  # already reported as broken
        stray = [f"{col}{row}" for row in range(first_row + 1, sheet.max_row + 1)
                 if is_formula(cell_value(sheet[f"{col}{row}"]))]
        if stray:
            report.add(
                SERIOUS, sheet_name, ", ".join(stray[:6]),
                f"Formula written inside the spill range of column {col}",
                f"{len(stray)} cell(s) below the anchor hold their own formula: "
                f"{', '.join(stray[:6])}. A spilled range should contain nothing "
                f"but the anchor.",
                f"Clear them and let {spec['spill_anchors'][col]} fill the column.",
            )

    payments = read_payments(sheet, spec)
    if not payments:
        report.add(SERIOUS, sheet_name, f"A{first_row}",
                   "Per-payment sheet is empty",
                   "No payment rows found, yet the sheet is configured as the "
                   "source for every dividend aggregate.",
                   "Check first_data_row in the config against the sheet.")
        return

    # --- payment type vocabulary ------------------------------------------- #
    type_col = column_by_meaning(columns, "payment_type", "I")
    allowed = [str(v).lower()
               for v in (columns.get(type_col, {}).get("allowed") or [])]
    if allowed:
        bad = [f"{type_col}{p['row']}={p['type'] or '(empty)'}"
               for p in payments if p["type"] not in allowed]
        if bad:
            report.add(
                SERIOUS, sheet_name, ", ".join(bad[:6]),
                f"Unrecognised payment type in {len(bad)} row(s)",
                f"Allowed: {', '.join(allowed)}. Found: {', '.join(bad[:6])}"
                f"{f' and {len(bad) - 6} more' if len(bad) > 6 else ''}. The "
                f"per-share rate in MyShares filters on type, so a typo silently "
                f"drops the row or lets a special dividend pose as the regular rate.",
                "Correct the value to one of the allowed strings.",
            )

    # --- dates -------------------------------------------------------------- #
    bad_dates = [f"A{p['row']}/B{p['row']}" for p in payments
                 if not p["year"] or not p["month"]
                 or not 1 <= p["month"] <= 12]
    if bad_dates:
        report.add(
            SERIOUS, sheet_name, ", ".join(bad_dates[:6]),
            f"Missing or impossible payment date in {len(bad_dates)} row(s)",
            f"{', '.join(bad_dates[:6])}. Year and month drive both the tax factor "
            f"lookup and the trailing-twelve-month window.",
            "Fill in the payment year and month, 1-12.",
        )

    # --- tickers must exist in the positions sheet -------------------------- #
    positions = next((name for name, s in config.get("sheets", {}).items()
                      if s.get("role") == "open_positions"
                      and name in workbook.sheetnames), None)
    if positions:
        pos_spec = config["sheets"][positions]
        pos_ticker_col = column_by_meaning(pos_spec.get("columns", {}), "ticker", "B")
        known = set()
        for block in ("dgi", "non_dgi"):
            known.update(tickers_in_rows(workbook[positions], pos_ticker_col,
                                         block_rows(pos_spec, block)).values())
        unknown = sorted({p["ticker"] for p in payments if p["ticker"] not in known})
        if unknown and known:
            report.add(
                WARNING, sheet_name, f"column {column_by_meaning(columns, 'ticker', 'C')}",
                f"{len(unknown)} ticker(s) not present in {positions}",
                f"{', '.join(unknown[:8])}. Either the position was sold - in which "
                f"case the history is correct and this is expected - or the ticker is "
                f"misspelled, and the SUMIFS that aggregates it silently returns zero.",
                f"Confirm each one is a closed position rather than a typo.",
            )


CHECKS = [
    check_dead_google_formulas,
    check_total_ranges,
    check_non_dgi_separator,
    check_non_dgi_isolation,
    check_weight_denominators,
    check_tax_factors,
    check_dividend_history_shape,
    check_hardcoded_frequency,
    check_labels,
    check_tickers,
    check_sectors,
    check_formula_errors,
    check_fragile_references,
    check_type_consistency,
    check_payment_sheet,
]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx")
    parser.add_argument("--config", default=None)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    config_path = Path(args.config) if args.config else \
        Path(__file__).resolve().parents[3] / "config" / "portfolio.json"
    config = json.loads(config_path.read_text(encoding="utf-8")) \
        if config_path.exists() else {}

    workbook = openpyxl.load_workbook(args.xlsx, data_only=False)
    report = Report()

    for check in CHECKS:
        try:
            check(workbook, config, report)
        except Exception as error:  # a failing check must not abort the audit
            report.add(INFO, "(audit)", "", f"Check {check.__name__} failed",
                       f"{type(error).__name__}: {error}", "Report this bug.")

    if args.json:
        print(json.dumps({"file": args.xlsx, "findings": report.sorted_findings()},
                         ensure_ascii=False, indent=2))
    else:
        print(format_text(report, args.xlsx))

    sys.exit(1 if report.count(CRITICAL) else 0)


if __name__ == "__main__":
    main()
